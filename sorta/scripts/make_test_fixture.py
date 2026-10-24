"""生成手工演示用的模拟下载目录（默认 tests/fixtures/downloads）。

自动化测试一律使用 pytest tmp_path 动态构造（开发文档 §13.2），
本脚本仅供 CLI 人工演示与调试使用。
"""

from __future__ import annotations

import argparse
import io
import json
import sys
from pathlib import Path

from docx import Document

ROOT = Path(__file__).resolve().parents[1]


def pdf_escape(text: str) -> str:
    return text.replace("\\", r"\\").replace("(", r"\(").replace(")", r"\)")


def build_pdf_bytes(text: str) -> bytes:
    content = f"BT /F1 18 Tf 72 720 Td ({pdf_escape(text)}) Tj ET".encode("latin-1")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >>",
        b"<< /Length " + str(len(content)).encode() + b" >>\nstream\n" + content + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets: list[int] = []
    for i, obj in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(f"{i} 0 obj ".encode() + obj + b" endobj\n")
    xref_pos = out.tell()
    out.write(f"xref\n0 {len(objects) + 1}\n".encode())
    out.write(b"0000000000 65535 f \n")
    for off in offsets:
        out.write(f"{off:010d} 00000 n \n".encode())
    trailer = f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref_pos}\n%%EOF\n"
    out.write(trailer.encode())
    return out.getvalue()


def build(out_dir: Path) -> None:
    out_dir.mkdir(parents=True, exist_ok=True)
    (out_dir / "web").mkdir(exist_ok=True)
    (out_dir / "deep" / "a" / "b" / "c").mkdir(parents=True, exist_ok=True)

    doc = Document()
    doc.add_paragraph("软件测试任务书：使用 Postman 完成接口测试并导出 Collection。")
    doc.add_paragraph("交付材料：collection.json 与运行截图。")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "课程"
    table.cell(0, 1).text = "软件测试"
    doc.save(out_dir / "任务书-软件测试.docx")

    (out_dir / "postman_collection.json").write_text(
        json.dumps({"info": {"name": "homework"}, "item": []}, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    (out_dir / "SQL注入学习笔记.txt").write_text(
        "SQL 注入学习笔记：\n- union select\n- 盲注\n- 堆叠注入\n", encoding="utf-8"
    )
    (out_dir / "逆向工程入门.pdf").write_bytes(build_pdf_bytes("Reverse engineering with IDA"))
    (out_dir / "old_gbk_notes.txt").write_text("GBK 编码的旧笔记", encoding="gbk")
    (out_dir / "screenshot_01.png").write_bytes(b"\x89PNG\r\n\x1a\n fake image bytes")
    (out_dir / "web" / "crawler.py").write_text("# 爬虫脚本\nprint('hi')\n", encoding="utf-8")
    (out_dir / "web" / "README.md").write_text("# Web 学习资料\n", encoding="utf-8")
    (out_dir / "deep" / "a" / "b" / "c" / "buried.txt").write_text("too deep", encoding="utf-8")
    (out_dir / ".hidden-note.txt").write_text("hidden", encoding="utf-8")

    print(f"已生成模拟下载目录: {out_dir}")


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--out",
        type=Path,
        default=ROOT / "tests" / "fixtures" / "downloads",
        help="输出目录（默认 tests/fixtures/downloads）",
    )
    args = parser.parse_args(argv)
    build(args.out)
    print(f'演示命令: uv run sorta scan "{args.out}"')
    return 0


if __name__ == "__main__":
    sys.exit(main())
