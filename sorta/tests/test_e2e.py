"""M1 端到端：模拟下载目录 → 扫描注册 → 按 file_id 读取（开发文档 §18 场景）。

不连接 LLM；验证"能扫描授权目录并安全读取支持的文件"这一 M1 完成标准。
"""

from __future__ import annotations

from pathlib import Path

from docx import Document

from sorta.config import SortaConfig
from sorta.scanner import FileRegistry, ReadStatus, Scanner, read_file


def _build_downloads(root: Path, make_pdf) -> None:
    (root / "web").mkdir(parents=True)
    (root / "deep" / "a" / "b" / "c").mkdir(parents=True)

    (root / "任务书-软件测试.docx").touch()
    doc = Document()
    doc.add_paragraph("软件测试任务书：使用 Postman 完成接口测试并导出 Collection。")
    doc.add_paragraph("交付材料：collection.json 与运行截图。")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "课程"
    table.cell(0, 1).text = "软件测试"
    doc.save(root / "任务书-软件测试.docx")

    (root / "postman_collection.json").write_text(
        '{"info": {"name": "homework"}, "item": []}', encoding="utf-8"
    )
    (root / "SQL注入学习笔记.txt").write_text(
        "SQL 注入：union select、盲注、堆叠注入。", encoding="utf-8"
    )
    make_pdf(root / "逆向工程入门.pdf", "Reverse engineering with IDA")
    (root / "old_gbk_notes.txt").write_text("GBK 编码的旧笔记", encoding="gbk")
    (root / "screenshot_01.png").write_bytes(b"\x89PNG\r\n\x1a\n fake image bytes")
    (root / "web" / "crawler.py").write_text("# 爬虫脚本\nprint('hi')\n", encoding="utf-8")
    (root / "deep" / "a" / "b" / "c" / "buried.txt").write_text("too deep", encoding="utf-8")


def test_end_to_end_scan_then_read(tmp_path, make_pdf):
    root = tmp_path / "downloads"
    root.mkdir()
    _build_downloads(root, make_pdf)

    registry = FileRegistry()
    scanner = Scanner(registry, SortaConfig())
    page = scanner.scan(root)

    # 深度 3 的文件不出现；其余普通文件全部注册
    paths = {f.relative_path for f in page.files}
    assert "deep/a/b/c/buried.txt" not in paths
    assert "任务书-软件测试.docx" in paths
    assert "postman_collection.json" in paths
    assert "web/crawler.py" in paths

    # 通过 file_id 读取（注册表是唯一出口，路径经过边界校验）
    by_path = {f.relative_path: f for f in page.files}
    limits = {"max_chars": 5_000, "max_bytes": 20 * 1024 * 1024}

    record = by_path["任务书-软件测试.docx"]
    _, path = registry.resolve_path(page.scan_id, record.id)
    r = read_file(path, file_id=record.id, **limits)
    assert r.status is ReadStatus.OK and "Postman" in r.content and "collection.json" in r.content

    record = by_path["postman_collection.json"]
    _, path = registry.resolve_path(page.scan_id, record.id)
    r = read_file(path, file_id=record.id, **limits)
    assert r.status is ReadStatus.OK and "homework" in r.content

    record = by_path["SQL注入学习笔记.txt"]
    _, path = registry.resolve_path(page.scan_id, record.id)
    r = read_file(path, file_id=record.id, **limits)
    assert r.status is ReadStatus.OK and "盲注" in r.content

    record = by_path["逆向工程入门.pdf"]
    _, path = registry.resolve_path(page.scan_id, record.id)
    r = read_file(path, file_id=record.id, **limits)
    assert r.status is ReadStatus.OK and "Reverse engineering" in r.content

    record = by_path["old_gbk_notes.txt"]
    _, path = registry.resolve_path(page.scan_id, record.id)
    r = read_file(path, file_id=record.id, **limits)
    assert r.status is ReadStatus.OK and "旧笔记" in r.content

    # 无关二进制文件：如实报告不可读，而不是猜测内容
    record = by_path["screenshot_01.png"]
    _, path = registry.resolve_path(page.scan_id, record.id)
    r = read_file(path, file_id=record.id, **limits)
    assert r.status is ReadStatus.UNSUPPORTED

    # 扫描后文件被删除：注册表如实报告 missing
    (root / "web" / "crawler.py").unlink()
    record = by_path["web/crawler.py"]
    _, path = registry.resolve_path(page.scan_id, record.id)
    r = read_file(path, file_id=record.id, **limits)
    assert r.status is ReadStatus.MISSING
