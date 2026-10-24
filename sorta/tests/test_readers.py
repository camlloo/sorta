"""Readers 测试（开发文档 §13.1：文本截断、编码异常、损坏 DOCX、
加密 PDF、扫描型 PDF、超大文件、不支持的格式）。"""

from __future__ import annotations

import io

import pytest
from docx import Document
from pypdf import PdfWriter

from sorta.scanner import ReadStatus, read_file

LIMITS = {"max_chars": 5_000, "max_bytes": 20 * 1024 * 1024}


def test_read_txt_utf8_chinese(tmp_path):
    p = tmp_path / "note.txt"
    p.write_text("SQL 注入学习笔记：union select", encoding="utf-8")
    r = read_file(p, **LIMITS)
    assert r.status is ReadStatus.OK
    assert "SQL 注入" in r.content
    assert r.truncated is False


@pytest.mark.parametrize("ext", ["md", "json", "csv", "py", "js", "ts"])
def test_read_supported_text_formats(tmp_path, ext):
    p = tmp_path / f"f.{ext}"
    p.write_text("hello world", encoding="utf-8")
    r = read_file(p, **LIMITS)
    assert r.status is ReadStatus.OK and r.content == "hello world"


def test_read_txt_gbk_fallback(tmp_path):
    p = tmp_path / "legacy.txt"
    p.write_text("老式编码的中文笔记", encoding="gbk")
    r = read_file(p, **LIMITS)
    assert r.status is ReadStatus.OK
    assert "老式编码的中文笔记" in r.content


def test_read_txt_truncation(tmp_path):
    p = tmp_path / "long.txt"
    p.write_text("A" * 100, encoding="utf-8")
    r = read_file(p, max_chars=10, max_bytes=1024)
    assert r.status is ReadStatus.OK
    assert r.content == "A" * 10 and r.truncated is True
    assert "截断" in r.error


def test_read_binary_rejected(tmp_path):
    p = tmp_path / "blob.txt"
    p.write_bytes(b"\x00\x01\x02binary\x00junk")
    r = read_file(p, **LIMITS)
    assert r.status is ReadStatus.UNREADABLE
    assert r.content is None


def test_read_undecodable_rejected(tmp_path):
    p = tmp_path / "weird.txt"
    # utf-8 与 gbk 都无法解码的字节序列（无 NUL，避开二进制快检）
    p.write_bytes(b"\xc3\x28\xff\xfe\xaa\xad\xef\xbb hello")
    r = read_file(p, **LIMITS)
    assert r.status is ReadStatus.UNREADABLE


def test_read_unsupported_format(tmp_path):
    p = tmp_path / "image.png"
    p.write_bytes(b"\x89PNG fake data")
    r = read_file(p, **LIMITS)
    assert r.status is ReadStatus.UNSUPPORTED and r.content is None


def test_read_too_large_rejected_before_parsing(tmp_path):
    p = tmp_path / "huge.txt"
    p.write_bytes(b"0123456789" * 10)  # 100 字节
    r = read_file(p, max_chars=5000, max_bytes=50)
    assert r.status is ReadStatus.TOO_LARGE
    assert "超过上限" in r.error


def test_read_missing_file(tmp_path):
    r = read_file(tmp_path / "ghost.txt", **LIMITS)
    assert r.status is ReadStatus.MISSING


def test_read_pdf_ok(make_pdf, tmp_path):
    p = tmp_path / "paper.pdf"
    make_pdf(p, "Reverse engineering basics")
    r = read_file(p, **LIMITS)
    assert r.status is ReadStatus.OK
    assert "Reverse engineering basics" in r.content


def test_read_pdf_encrypted(make_encrypted_pdf, tmp_path):
    p = tmp_path / "locked.pdf"
    make_encrypted_pdf(p)
    r = read_file(p, **LIMITS)
    assert r.status is ReadStatus.ENCRYPTED and r.content is None
    assert "加密" in r.error


def test_read_pdf_scanned_no_text(tmp_path):
    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)  # 空白页 = 扫描型 PDF 的等价物
    buf = io.BytesIO()
    writer.write(buf)
    p = tmp_path / "scanned.pdf"
    p.write_bytes(buf.getvalue())

    r = read_file(p, **LIMITS)
    assert r.status is ReadStatus.NO_TEXT
    assert "扫描型" in r.error


def test_read_pdf_corrupt(tmp_path):
    p = tmp_path / "broken.pdf"
    p.write_bytes(b"%PDF-1.7 this is not really a pdf \x00\x01 garbage")
    r = read_file(p, **LIMITS)
    assert r.status is ReadStatus.CORRUPT


def test_read_docx_ok_with_table(tmp_path):
    doc = Document()
    doc.add_paragraph("软件测试课程任务书")
    doc.add_paragraph("要求：导出 Postman Collection 并截图")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "交付物"
    table.cell(0, 1).text = "collection.json"
    p = tmp_path / "任务书.docx"
    doc.save(p)

    r = read_file(p, **LIMITS)
    assert r.status is ReadStatus.OK
    assert "任务书" in r.content and "collection.json" in r.content


def test_read_docx_corrupt(tmp_path):
    p = tmp_path / "broken.docx"
    p.write_bytes(b"PK\x03\x04 definitely not a valid zip archive payload")
    r = read_file(p, **LIMITS)
    assert r.status is ReadStatus.CORRUPT
