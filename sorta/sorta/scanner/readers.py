"""受控文件内容读取（开发文档 §6.3）。

只提取文本，绝不执行文件内容中的任何指令——文件内容一律视为不可信数据
（§6.5 / §10.2 Prompt Injection 防护的第一道边界）。

支持：txt / md / json / csv / py / js / ts（文本，utf-8 优先、gbk 回退）、
pdf（pypdf，不做 OCR）、docx（python-docx，段落 + 表格）。
所有异常状态显式返回，不抛出、不猜测内容。
"""

from __future__ import annotations

import os
from enum import StrEnum
from pathlib import Path

from pydantic import BaseModel

TEXT_EXTENSIONS = frozenset({"txt", "md", "json", "csv", "py", "js", "ts"})
_TEXT_ENCODINGS = ("utf-8-sig", "gbk")  # utf-8-sig 兼容带/不带 BOM 的 UTF-8
_MAX_PDF_PAGES = 200  # 页数上限：防资源耗尽（§10.2）


class ReadStatus(StrEnum):
    OK = "ok"
    UNSUPPORTED = "unsupported_format"
    TOO_LARGE = "too_large"
    MISSING = "missing"
    UNREADABLE = "unreadable"  # 二进制/无法解码
    ENCRYPTED = "encrypted"
    NO_TEXT = "no_text"  # 合法文档但无可提取文本（如扫描型 PDF）
    CORRUPT = "corrupt"  # 格式损坏，解析器报错


class ReadResult(BaseModel):
    status: ReadStatus
    format: str
    file_id: str | None = None
    content: str | None = None
    truncated: bool = False
    error: str | None = None


def read_file(
    path: str | Path,
    *,
    max_chars: int,
    max_bytes: int,
    file_id: str | None = None,
) -> ReadResult:
    """读取单个文件文本。先做存在性与大小检查，再按扩展名分派解析器。"""
    path = Path(path)
    ext = path.suffix.lower().lstrip(".")
    try:
        st = os.stat(path)
    except (FileNotFoundError, NotADirectoryError):
        return ReadResult(
            status=ReadStatus.MISSING, format=ext, file_id=file_id, error="文件不存在"
        )
    except OSError as exc:
        return ReadResult(status=ReadStatus.UNREADABLE, format=ext, file_id=file_id, error=str(exc))
    if st.st_size > max_bytes:
        return ReadResult(
            status=ReadStatus.TOO_LARGE,
            format=ext,
            file_id=file_id,
            error=f"文件 {st.st_size} 字节超过上限 {max_bytes}",
        )
    try:
        if ext in TEXT_EXTENSIONS:
            return _read_text(path, ext, file_id, max_chars)
        if ext == "pdf":
            return _read_pdf(path, file_id, max_chars)
        if ext == "docx":
            return _read_docx(path, file_id, max_chars)
    except Exception as exc:  # 解析器对损坏文件抛出的各类异常统一归为 corrupt
        return ReadResult(
            status=ReadStatus.CORRUPT,
            format=ext,
            file_id=file_id,
            error=f"{type(exc).__name__}: {exc}",
        )
    return ReadResult(
        status=ReadStatus.UNSUPPORTED, format=ext, file_id=file_id, error=f"不支持的格式: .{ext}"
    )


def _truncate(text: str, max_chars: int) -> tuple[str, bool]:
    if len(text) <= max_chars:
        return text, False
    return text[:max_chars], True


def _looks_binary(text: str) -> bool:
    sample = text[:2000]
    if not sample:
        return False
    ctrl = sum(1 for ch in sample if ord(ch) < 32 and ch not in "\t\n\r\f\v")
    return ctrl / len(sample) > 0.02


def _read_text(path: Path, ext: str, file_id: str | None, max_chars: int) -> ReadResult:
    raw = path.read_bytes()
    if b"\x00" in raw[:8192]:
        return ReadResult(
            status=ReadStatus.UNREADABLE,
            format=ext,
            file_id=file_id,
            error="内容包含 NUL 字节，疑似二进制文件",
        )
    text: str | None = None
    used_encoding = ""
    for encoding in _TEXT_ENCODINGS:
        try:
            text = raw.decode(encoding)
            used_encoding = encoding
            break
        except UnicodeDecodeError:
            continue
    if text is None:
        return ReadResult(
            status=ReadStatus.UNREADABLE,
            format=ext,
            file_id=file_id,
            error="无法以 utf-8/gbk 解码",
        )
    if _looks_binary(text):
        return ReadResult(
            status=ReadStatus.UNREADABLE,
            format=ext,
            file_id=file_id,
            error="解码成功但控制字符比例异常，疑似二进制",
        )
    content, truncated = _truncate(text, max_chars)
    return ReadResult(
        status=ReadStatus.OK,
        format=ext,
        file_id=file_id,
        content=content,
        truncated=truncated,
        error=None if not truncated else f"已截断至 {max_chars} 字符（编码 {used_encoding}）",
    )


def _read_pdf(path: Path, file_id: str | None, max_chars: int) -> ReadResult:
    from pypdf import PdfReader

    reader = PdfReader(str(path), strict=False)
    if reader.is_encrypted:
        return ReadResult(
            status=ReadStatus.ENCRYPTED,
            format="pdf",
            file_id=file_id,
            error="PDF 已加密，未提供密码",
        )
    parts: list[str] = []
    total = 0
    for i, page in enumerate(reader.pages):
        if i >= _MAX_PDF_PAGES or total >= max_chars * 2:
            break
        try:
            page_text = page.extract_text() or ""
        except Exception:  # 单页解析失败不致命，跳过该页继续
            page_text = ""
        parts.append(page_text)
        total += len(page_text)
    text = "\n".join(parts)
    if not text.strip():
        return ReadResult(
            status=ReadStatus.NO_TEXT,
            format="pdf",
            file_id=file_id,
            error="无法提取文本（可能是扫描型/纯图片 PDF，不做 OCR）",
        )
    content, truncated = _truncate(text, max_chars)
    return ReadResult(
        status=ReadStatus.OK,
        format="pdf",
        file_id=file_id,
        content=content,
        truncated=truncated,
        error=None if not truncated else f"已截断至 {max_chars} 字符",
    )


def _read_docx(path: Path, file_id: str | None, max_chars: int) -> ReadResult:
    from docx import Document

    doc = Document(str(path))
    chunks = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            chunks.extend(cell.text for cell in row.cells)
    text = "\n".join(chunks)
    if not text.strip():
        return ReadResult(
            status=ReadStatus.NO_TEXT,
            format="docx",
            file_id=file_id,
            error="文档中没有可提取的文本（可能仅含图片）",
        )
    content, truncated = _truncate(text, max_chars)
    return ReadResult(
        status=ReadStatus.OK,
        format="docx",
        file_id=file_id,
        content=content,
        truncated=truncated,
        error=None if not truncated else f"已截断至 {max_chars} 字符",
    )
